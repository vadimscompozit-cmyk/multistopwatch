"""Секундомер, который сразу пишет результаты в Excel-таблицу.

Файл .xlsx создаётся в момент первого старта и переписывается после
каждого круга, поэтому данные не теряются даже при аварийном закрытии.
"""

import os
import sys
import time
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
RECORDS_DIR = BASE_DIR / "записи"

HEADERS = [
    "№",
    "Время круга",
    "Общее время",
    "Круг, сек",
    "Всего, сек",
    "Время суток",
    "Комментарий",
]
COLUMN_WIDTHS = [6, 14, 14, 12, 12, 14, 40]

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
SUMMARY_FILL = PatternFill("solid", fgColor="E8EEF7")
THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def format_time(seconds: float) -> str:
    """0:07.35 -> '00:07.35', больше часа -> '1:02:03.45'."""
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    if hours >= 1:
        return f"{int(hours)}:{int(minutes):02d}:{secs:05.2f}"
    return f"{int(minutes):02d}:{secs:05.2f}"


class Stopwatch(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Секундомер → Excel")
        self.minsize(640, 520)
        self.configure(padx=16, pady=14)

        self.running = False
        self.elapsed = 0.0          # накоплено на момент последней паузы
        self.started_at = None      # perf_counter() текущего запуска
        self.laps = []              # [{"split", "total", "clock", "note"}]
        self.excel_path = None
        self.save_failed = False

        self._build_ui()
        self._bind_keys()
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        self._tick()

    # ------------------------------------------------------------- интерфейс
    def _build_ui(self):
        style = ttk.Style(self)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        style.configure("Big.TButton", font=("Segoe UI", 11), padding=(10, 8))

        self.display = tk.Label(
            self,
            text="00:00.00",
            font=("Consolas", 54, "bold"),
            fg="#1F3864",
        )
        self.display.pack(pady=(4, 2))

        self.file_label = tk.Label(
            self,
            text="Таблица будет создана при первом старте",
            font=("Segoe UI", 9),
            fg="#666666",
        )
        self.file_label.pack()

        buttons = ttk.Frame(self)
        buttons.pack(fill="x", pady=12)
        for i in range(4):
            buttons.columnconfigure(i, weight=1)

        self.start_btn = ttk.Button(
            buttons, text="Старт  (Пробел)", style="Big.TButton",
            command=self.toggle, takefocus=False,
        )
        self.start_btn.grid(row=0, column=0, sticky="ew", padx=3)

        self.lap_btn = ttk.Button(
            buttons, text="Круг  (Enter)", style="Big.TButton",
            command=self.add_lap, state="disabled", takefocus=False,
        )
        self.lap_btn.grid(row=0, column=1, sticky="ew", padx=3)

        self.reset_btn = ttk.Button(
            buttons, text="Сброс  (Ctrl+R)", style="Big.TButton",
            command=self.reset, takefocus=False,
        )
        self.reset_btn.grid(row=0, column=2, sticky="ew", padx=3)

        self.open_btn = ttk.Button(
            buttons, text="Открыть таблицу", style="Big.TButton",
            command=self.open_file, state="disabled", takefocus=False,
        )
        self.open_btn.grid(row=0, column=3, sticky="ew", padx=3)

        note_row = ttk.Frame(self)
        note_row.pack(fill="x", pady=(0, 10))
        ttk.Label(note_row, text="Комментарий к кругу:").pack(side="left")
        self.note_var = tk.StringVar()
        self.note_entry = ttk.Entry(note_row, textvariable=self.note_var)
        self.note_entry.pack(side="left", fill="x", expand=True, padx=(8, 0))
        self.note_entry.bind("<Return>", lambda _e: self.add_lap())

        columns = ("num", "split", "total", "clock", "note")
        self.tree = ttk.Treeview(self, columns=columns, show="headings", height=10)
        for col, title, width, anchor in (
            ("num", "№", 50, "center"),
            ("split", "Круг", 110, "center"),
            ("total", "Всего", 110, "center"),
            ("clock", "Время", 90, "center"),
            ("note", "Комментарий", 220, "w"),
        ):
            self.tree.heading(col, text=title)
            self.tree.column(col, width=width, anchor=anchor)
        self.tree.pack(fill="both", expand=True)

        self.status = tk.Label(self, text="", font=("Segoe UI", 9), fg="#666666",
                               anchor="w")
        self.status.pack(fill="x", pady=(8, 0))

    def _bind_keys(self):
        self.bind("<space>", self._on_space)
        self.bind("<Control-r>", lambda _e: self.reset())
        self.bind("<Control-R>", lambda _e: self.reset())

    def _on_space(self, _event):
        if self.focus_get() is self.note_entry:
            return None
        self.toggle()
        return "break"

    # -------------------------------------------------------------- механика
    def current_elapsed(self) -> float:
        if self.running:
            return self.elapsed + (time.perf_counter() - self.started_at)
        return self.elapsed

    def _tick(self):
        self.display.config(text=format_time(self.current_elapsed()))
        if self.save_failed:
            self.save_workbook(quiet=True)
        self.after(30, self._tick)

    def toggle(self):
        if self.running:
            self.elapsed = self.current_elapsed()
            self.running = False
            self.started_at = None
            self.start_btn.config(text="Продолжить  (Пробел)")
            self.set_status("Пауза")
        else:
            if self.excel_path is None:
                self.create_workbook()
            self.started_at = time.perf_counter()
            self.running = True
            self.start_btn.config(text="Пауза  (Пробел)")
            self.lap_btn.config(state="normal")
            self.set_status("Идёт отсчёт")

    def add_lap(self):
        if not self.running and not self.laps and self.elapsed == 0:
            return
        total = self.current_elapsed()
        previous = self.laps[-1]["total"] if self.laps else 0.0
        lap = {
            "split": total - previous,
            "total": total,
            "clock": datetime.now().strftime("%H:%M:%S"),
            "note": self.note_var.get().strip(),
        }
        self.laps.append(lap)
        self.note_var.set("")

        self.tree.insert(
            "", "end",
            values=(len(self.laps), format_time(lap["split"]),
                    format_time(lap["total"]), lap["clock"], lap["note"]),
        )
        self.tree.yview_moveto(1.0)
        self.save_workbook()

    def reset(self):
        if self.laps and not messagebox.askyesno(
            "Сброс",
            f"Записано кругов: {len(self.laps)}.\n"
            f"Таблица уже сохранена — начать новый замер в новом файле?",
        ):
            return
        self.running = False
        self.started_at = None
        self.elapsed = 0.0
        self.laps.clear()
        self.excel_path = None
        self.save_failed = False
        self.tree.delete(*self.tree.get_children())
        self.start_btn.config(text="Старт  (Пробел)")
        self.lap_btn.config(state="disabled")
        self.open_btn.config(state="disabled")
        self.file_label.config(text="Таблица будет создана при первом старте")
        self.set_status("Сброшено — новый файл создастся при старте")

    def set_status(self, text: str, error: bool = False):
        self.status.config(text=text, fg="#B00020" if error else "#666666")

    # ----------------------------------------------------------------- Excel
    def create_workbook(self):
        RECORDS_DIR.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y-%m-%d %H-%M-%S")
        self.excel_path = RECORDS_DIR / f"секундомер {stamp}.xlsx"
        self.save_workbook()
        self.open_btn.config(state="normal")
        self.file_label.config(text=str(self.excel_path))

    def save_workbook(self, quiet: bool = False):
        if self.excel_path is None:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Замер"

        ws["A1"] = f"Секундомер — {datetime.now().strftime('%d.%m.%Y')}"
        ws["A1"].font = Font(size=14, bold=True, color="1F3864")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

        for idx, (header, width) in enumerate(zip(HEADERS, COLUMN_WIDTHS), start=1):
            cell = ws.cell(row=3, column=idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.row_dimensions[3].height = 22
        ws.freeze_panes = "A4"

        row = 4
        for number, lap in enumerate(self.laps, start=1):
            values = [
                number,
                format_time(lap["split"]),
                format_time(lap["total"]),
                round(lap["split"], 2),
                round(lap["total"], 2),
                lap["clock"],
                lap["note"],
            ]
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=column, value=value)
                cell.border = BORDER
                if column in (4, 5):
                    cell.number_format = "0.00"
                if column != 7:
                    cell.alignment = Alignment(horizontal="center")
            row += 1

        splits = [lap["split"] for lap in self.laps]
        summary = [
            ("Кругов", len(self.laps)),
            ("Общее время", format_time(self.laps[-1]["total"]) if self.laps else "—"),
            ("Средний круг", format_time(sum(splits) / len(splits)) if splits else "—"),
            ("Самый быстрый", format_time(min(splits)) if splits else "—"),
            ("Самый медленный", format_time(max(splits)) if splits else "—"),
        ]
        row += 1
        for title, value in summary:
            label = ws.cell(row=row, column=1, value=title)
            label.font = Font(bold=True)
            label.fill = SUMMARY_FILL
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            result = ws.cell(row=row, column=3, value=value)
            result.fill = SUMMARY_FILL
            result.alignment = Alignment(horizontal="center")
            row += 1

        try:
            wb.save(self.excel_path)
        except PermissionError:
            self.save_failed = True
            self.set_status(
                f"Файл открыт в Excel — закройте его, запись возобновится сама "
                f"({len(self.laps)} кругов в памяти)",
                error=True,
            )
        except OSError as exc:
            self.save_failed = True
            self.set_status(f"Не удалось сохранить: {exc}", error=True)
        else:
            if self.save_failed:
                self.save_failed = False
                self.set_status("Запись в файл восстановлена")
            elif not quiet:
                self.set_status(f"Сохранено: {self.excel_path.name}")

    def open_file(self):
        if self.excel_path and self.excel_path.exists():
            os.startfile(self.excel_path)

    def on_close(self):
        if self.laps:
            self.save_workbook(quiet=True)
        self.destroy()


if __name__ == "__main__":
    if sys.platform == "win32":
        try:  # чёткий шрифт на мониторах с масштабированием
            from ctypes import windll

            windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    Stopwatch().mainloop()
